"""``PresentationManager``: the single source of truth for the open deck.

All tool handlers operate on one shared ``PresentationManager`` instance
instead of touching ``python-pptx`` objects directly. It owns the
currently open ``Presentation``, the file path it was loaded from/will be
saved to, and which slide is "active" (so the LLM can say "add a bullet
list" without repeating a slide index every time).
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.slide import Slide

from ppt_agent.exceptions import ToolExecutionError

# Index of each named layout in python-pptx's default template. Layout
# order is a stable property of the built-in template python-pptx ships
# with; a custom template passed to `new()` may not match these names.
LAYOUT_NAME_TO_INDEX = {
    "title": 0,
    "title_and_content": 1,
    "section_header": 2,
    "two_content": 3,
    "comparison": 4,
    "title_only": 5,
    "blank": 6,
}


class PresentationManager:
    """Owns the currently open ``python-pptx`` presentation and its state."""

    def __init__(self) -> None:
        self._presentation: Presentation | None = None
        self._path: Path | None = None
        self._active_slide_index: int | None = None

    # -- lifecycle -----------------------------------------------------

    def new(self, template_path: str | None = None) -> None:
        """Start a fresh presentation, optionally based on a .pptx template.

        The new deck has no associated save path yet, even when built from
        a template file, since saving to the template's own path would
        overwrite it.
        """
        self._presentation = (
            Presentation(template_path) if template_path else Presentation()
        )
        self._path = None
        self._active_slide_index = None

    def open(self, path: str) -> None:
        """Load an existing .pptx file for editing."""
        resolved_path = Path(path).expanduser().resolve()
        if not resolved_path.exists():
            raise ToolExecutionError(f"Presentation file not found: {resolved_path}")
        self._presentation = Presentation(str(resolved_path))
        self._path = resolved_path
        slide_count = len(list(self._presentation.slides))
        self._active_slide_index = slide_count - 1 if slide_count > 0 else None

    def save(self, path: str | None = None) -> Path:
        """Save the current presentation, defaulting to the path it was opened/saved from."""
        self._require_presentation()
        if path is not None:
            self._path = Path(path).expanduser().resolve()
        if self._path is None:
            raise ToolExecutionError(
                "No file path is known for this presentation yet. "
                "Call save with an explicit path, e.g. save_presentation(path='deck.pptx')."
            )
        self._path.parent.mkdir(parents=True, exist_ok=True)
        self._presentation.save(str(self._path))
        return self._path

    @property
    def is_open(self) -> bool:
        return self._presentation is not None

    @property
    def path(self) -> Path | None:
        return self._path

    # -- slide access ----------------------------------------------------

    def add_slide(
        self, layout: str = "title_and_content", title: str | None = None
    ) -> int:
        """Append a new slide using a named layout and return its index."""
        self._require_presentation()
        layout_index = LAYOUT_NAME_TO_INDEX.get(layout)
        if layout_index is None:
            raise ToolExecutionError(
                f"Unknown layout '{layout}'. Supported layouts: {sorted(LAYOUT_NAME_TO_INDEX)}."
            )
        try:
            slide_layout = self._presentation.slide_layouts[layout_index]
        except IndexError as exc:
            raise ToolExecutionError(
                f"The current template has no layout at index {layout_index} for '{layout}'."
            ) from exc

        slide = self._presentation.slides.add_slide(slide_layout)
        if title is not None and slide.shapes.title is not None:
            slide.shapes.title.text = title

        new_index = len(self._presentation.slides) - 1
        self._active_slide_index = new_index
        return new_index

    def delete_slide(self, slide_index: int) -> None:
        """Remove a slide by index."""
        self._require_presentation()
        slide_id_list = self._presentation.slides._sldIdLst
        slides = list(self._presentation.slides)
        if not 0 <= slide_index < len(slides):
            raise ToolExecutionError(self._out_of_range_message(slide_index))
        slide_id = slide_id_list[slide_index]
        slide_id_list.remove(slide_id)
        if self._active_slide_index is not None and self._active_slide_index >= len(
            list(self._presentation.slides)
        ):
            self._active_slide_index = len(list(self._presentation.slides)) - 1 or None

    def get_slide(self, slide_index: int | None = None) -> Slide:
        """Resolve a slide index (or the active slide, if None) to a ``Slide``."""
        resolved_index = self.resolve_slide_index(slide_index)
        return list(self._presentation.slides)[resolved_index]

    def resolve_slide_index(self, slide_index: int | None) -> int:
        """Validate an explicit index, or fall back to the active slide."""
        self._require_presentation()
        if slide_index is None:
            if self._active_slide_index is None:
                raise ToolExecutionError(
                    "No slide is currently active and no slide_index was given. "
                    "Add a slide first, or pass an explicit slide_index."
                )
            return self._active_slide_index

        slide_count = len(list(self._presentation.slides))
        if not 0 <= slide_index < slide_count:
            raise ToolExecutionError(self._out_of_range_message(slide_index))
        self._active_slide_index = slide_index
        return slide_index

    def list_slides(self) -> list[dict]:
        """Summarize every slide for display or for the LLM's own bookkeeping."""
        self._require_presentation()
        summaries = []
        for index, slide in enumerate(self._presentation.slides):
            title_shape = slide.shapes.title
            summaries.append(
                {
                    "index": index,
                    "title": title_shape.text
                    if title_shape is not None and title_shape.has_text_frame
                    else "",
                    "shape_count": len(slide.shapes),
                    "is_active": index == self._active_slide_index,
                }
            )
        return summaries

    # -- internals ---------------------------------------------------------

    def _require_presentation(self) -> None:
        if self._presentation is None:
            raise ToolExecutionError(
                "No presentation is open yet. Create one first with create_presentation, "
                "or load an existing file with open_presentation."
            )

    def _out_of_range_message(self, slide_index: int) -> str:
        slide_count = len(list(self._presentation.slides))
        return f"slide_index {slide_index} is out of range (presentation has {slide_count} slide(s), 0-indexed)."
