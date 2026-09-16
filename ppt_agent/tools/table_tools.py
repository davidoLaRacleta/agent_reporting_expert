"""Tools for adding data tables to a slide."""

from __future__ import annotations

from langchain_core.tools import StructuredTool
from pptx.util import Inches
from pydantic import BaseModel, Field

from ppt_agent.exceptions import ToolExecutionError
from ppt_agent.presentation.manager import PresentationManager


class AddTableInput(BaseModel):
    rows: list[list[str]] = Field(
        description="Row-major 2D list of cell text. Every row must have the same length."
    )
    slide_index: int | None = Field(default=None, description="Defaults to the active slide.")
    left: float = Field(default=1.0, description="Left position in inches.")
    top: float = Field(default=1.5, description="Top position in inches.")
    width: float = Field(default=8.0, description="Width in inches.")
    height: float = Field(default=4.0, description="Height in inches.")
    header_row: bool = Field(default=True, description="Bold the first row as a header.")


def build_tools(manager: PresentationManager) -> list[StructuredTool]:
    """Build table-related tools bound to ``manager``."""

    def add_table(
        rows: list[list[str]],
        slide_index: int | None = None,
        left: float = 1.0,
        top: float = 1.5,
        width: float = 8.0,
        height: float = 4.0,
        header_row: bool = True,
    ) -> str:
        if not rows or not rows[0]:
            raise ToolExecutionError("rows must be a non-empty 2D list, e.g. [['A', 'B'], ['1', '2']].")

        column_count = len(rows[0])
        for row_index, row in enumerate(rows):
            if len(row) != column_count:
                raise ToolExecutionError(
                    f"Row {row_index} has {len(row)} cell(s) but row 0 has {column_count}; "
                    "every row must have the same number of columns."
                )

        slide = manager.get_slide(slide_index)
        graphic_frame = slide.shapes.add_table(
            len(rows), column_count, Inches(left), Inches(top), Inches(width), Inches(height)
        )
        table = graphic_frame.table
        for row_index, row_values in enumerate(rows):
            for column_index, cell_value in enumerate(row_values):
                cell = table.cell(row_index, column_index)
                cell.text = str(cell_value)
                if header_row and row_index == 0:
                    cell.text_frame.paragraphs[0].font.bold = True

        return f"Added a {len(rows)}x{column_count} table to slide {manager.resolve_slide_index(slide_index)}."

    return [
        StructuredTool.from_function(
            func=add_table,
            name="add_table",
            description="Add a table to a slide from a 2D list of cell text, with an optional bold header row.",
            args_schema=AddTableInput,
        ),
    ]
