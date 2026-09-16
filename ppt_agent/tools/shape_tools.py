"""Tools for adding and styling geometric shapes on a slide."""

from __future__ import annotations

from langchain_core.tools import StructuredTool
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from pydantic import BaseModel, Field

from ppt_agent.exceptions import ToolExecutionError
from ppt_agent.presentation.manager import PresentationManager
from ppt_agent.presentation.style_utils import hex_to_rgb_color

_SHAPE_TYPES = {
    "rectangle": MSO_SHAPE.RECTANGLE,
    "rounded_rectangle": MSO_SHAPE.ROUNDED_RECTANGLE,
    "oval": MSO_SHAPE.OVAL,
    "diamond": MSO_SHAPE.DIAMOND,
    "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
    "right_arrow": MSO_SHAPE.RIGHT_ARROW,
    "left_arrow": MSO_SHAPE.LEFT_ARROW,
    "up_arrow": MSO_SHAPE.UP_ARROW,
    "down_arrow": MSO_SHAPE.DOWN_ARROW,
    "pentagon": MSO_SHAPE.PENTAGON,
    "hexagon": MSO_SHAPE.HEXAGON,
    "chevron": MSO_SHAPE.CHEVRON,
    "star_5": MSO_SHAPE.STAR_5_POINT,
    "cloud": MSO_SHAPE.CLOUD,
}


class AddShapeInput(BaseModel):
    shape_type: str = Field(description=f"Supported: {sorted(_SHAPE_TYPES)}.")
    slide_index: int | None = Field(default=None, description="Defaults to the active slide.")
    left: float = Field(default=1.0, description="Left position in inches.")
    top: float = Field(default=1.0, description="Top position in inches.")
    width: float = Field(default=2.0, description="Width in inches.")
    height: float = Field(default=1.0, description="Height in inches.")
    text: str | None = Field(default=None, description="Optional text rendered inside the shape.")
    fill_color: str | None = Field(default=None, description="6-digit hex fill color, e.g. '4472C4'.")
    line_color: str | None = Field(default=None, description="6-digit hex outline color.")


class StyleShapeInput(BaseModel):
    shape_index: int = Field(description="Index of the shape within the slide's shape list.")
    slide_index: int | None = Field(default=None, description="Defaults to the active slide.")
    fill_color: str | None = Field(default=None, description="6-digit hex fill color.")
    line_color: str | None = Field(default=None, description="6-digit hex outline color.")
    line_width_pt: float | None = Field(default=None, description="Outline width in points.")


def build_tools(manager: PresentationManager) -> list[StructuredTool]:
    """Build shape-related tools bound to ``manager``."""

    def add_shape(
        shape_type: str,
        slide_index: int | None = None,
        left: float = 1.0,
        top: float = 1.0,
        width: float = 2.0,
        height: float = 1.0,
        text: str | None = None,
        fill_color: str | None = None,
        line_color: str | None = None,
    ) -> str:
        slide = manager.get_slide(slide_index)
        mso_shape_type = _SHAPE_TYPES.get(shape_type.lower())
        if mso_shape_type is None:
            raise ToolExecutionError(f"Unknown shape_type '{shape_type}'. Supported: {sorted(_SHAPE_TYPES)}.")

        shape = slide.shapes.add_shape(mso_shape_type, Inches(left), Inches(top), Inches(width), Inches(height))
        if fill_color is not None:
            shape.fill.solid()
            shape.fill.fore_color.rgb = hex_to_rgb_color(fill_color)
        if line_color is not None:
            shape.line.color.rgb = hex_to_rgb_color(line_color)
        if text is not None:
            shape.text_frame.text = text

        shape_index = len(slide.shapes) - 1
        return (
            f"Added '{shape_type}' shape (shape_index {shape_index}) to "
            f"slide {manager.resolve_slide_index(slide_index)}."
        )

    def style_shape(
        shape_index: int,
        slide_index: int | None = None,
        fill_color: str | None = None,
        line_color: str | None = None,
        line_width_pt: float | None = None,
    ) -> str:
        slide = manager.get_slide(slide_index)
        shape = _resolve_shape(slide, shape_index)

        if fill_color is not None:
            shape.fill.solid()
            shape.fill.fore_color.rgb = hex_to_rgb_color(fill_color)
        if line_color is not None:
            shape.line.color.rgb = hex_to_rgb_color(line_color)
        if line_width_pt is not None:
            shape.line.width = Pt(line_width_pt)

        return f"Updated style of shape {shape_index} on slide {manager.resolve_slide_index(slide_index)}."

    return [
        StructuredTool.from_function(
            func=add_shape,
            name="add_shape",
            description=f"Add a geometric shape to a slide. Supported shape_type values: {sorted(_SHAPE_TYPES)}.",
            args_schema=AddShapeInput,
        ),
        StructuredTool.from_function(
            func=style_shape,
            name="style_shape",
            description="Change the fill color, line color, and/or line width of an existing shape.",
            args_schema=StyleShapeInput,
        ),
    ]


def _resolve_shape(slide, shape_index: int):
    shapes = list(slide.shapes)
    if not 0 <= shape_index < len(shapes):
        raise ToolExecutionError(
            f"shape_index {shape_index} is out of range (slide has {len(shapes)} shape(s), 0-indexed)."
        )
    return shapes[shape_index]
