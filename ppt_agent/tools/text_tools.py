"""Tools for adding and styling text on a slide."""

from __future__ import annotations

from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from ppt_agent.exceptions import ToolExecutionError
from ppt_agent.presentation.manager import PresentationManager
from ppt_agent.presentation.style_utils import hex_to_rgb_color
from ppt_agent.tools.registry import ToolRegistry

_ALIGNMENTS = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
    "justify": PP_ALIGN.JUSTIFY,
}


def register(registry: ToolRegistry, manager: PresentationManager) -> None:
    """Register text-related tools onto ``registry``."""

    def set_slide_title(title: str, slide_index: int | None = None) -> str:
        slide = manager.get_slide(slide_index)
        if slide.shapes.title is None:
            raise ToolExecutionError(
                "This slide's layout has no title placeholder. Use add_text_box instead."
            )
        slide.shapes.title.text = title
        return f"Set title of slide {manager.resolve_slide_index(slide_index)} to '{title}'."

    registry.add(
        name="set_slide_title",
        description="Set the title placeholder text of a slide.",
        parameters={
            "type": "object",
            "properties": {
                "title": {"type": "string"},
                "slide_index": {
                    "type": "integer",
                    "description": "Defaults to the active slide.",
                },
            },
            "required": ["title"],
        },
        handler=set_slide_title,
    )

    def add_text_box(
        text: str,
        slide_index: int | None = None,
        left: float = 0.5,
        top: float = 0.5,
        width: float = 9.0,
        height: float = 1.0,
        font_size: int = 18,
        bold: bool = False,
        italic: bool = False,
        color: str | None = None,
        alignment: str = "left",
    ) -> str:
        slide = manager.get_slide(slide_index)
        text_box = slide.shapes.add_textbox(
            Inches(left), Inches(top), Inches(width), Inches(height)
        )
        text_frame = text_box.text_frame
        text_frame.word_wrap = True
        paragraph = text_frame.paragraphs[0]
        paragraph.text = text
        paragraph.alignment = _resolve_alignment(alignment)
        run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.italic = italic
        if color is not None:
            run.font.color.rgb = hex_to_rgb_color(color)
        return f"Added text box to slide {manager.resolve_slide_index(slide_index)}."

    registry.add(
        name="add_text_box",
        description="Add a free-floating text box with a single block of styled text to a slide.",
        parameters={
            "type": "object",
            "properties": {
                "text": {"type": "string"},
                "slide_index": {
                    "type": "integer",
                    "description": "Defaults to the active slide.",
                },
                "left": {"type": "number", "description": "Left position in inches."},
                "top": {"type": "number", "description": "Top position in inches."},
                "width": {"type": "number", "description": "Width in inches."},
                "height": {"type": "number", "description": "Height in inches."},
                "font_size": {"type": "integer", "description": "Font size in points."},
                "bold": {"type": "boolean"},
                "italic": {"type": "boolean"},
                "color": {
                    "type": "string",
                    "description": "6-digit hex color, e.g. '1F4E79'.",
                },
                "alignment": {"type": "string", "enum": sorted(_ALIGNMENTS)},
            },
            "required": ["text"],
        },
        handler=add_text_box,
    )

    def add_bullet_list(
        items: list[str],
        slide_index: int | None = None,
        left: float | None = None,
        top: float | None = None,
        width: float | None = None,
        height: float | None = None,
        font_size: int = 18,
    ) -> str:
        if not items:
            raise ToolExecutionError("items must be a non-empty list of strings.")

        slide = manager.get_slide(slide_index)
        text_frame = _resolve_body_text_frame(slide, left, top, width, height)
        text_frame.clear()
        text_frame.word_wrap = True
        for position, item_text in enumerate(items):
            paragraph = (
                text_frame.paragraphs[0]
                if position == 0
                else text_frame.add_paragraph()
            )
            paragraph.text = item_text
            paragraph.level = 0
            paragraph.font.size = Pt(font_size)
        return f"Added {len(items)} bullet point(s) to slide {manager.resolve_slide_index(slide_index)}."

    registry.add(
        name="add_bullet_list",
        description=(
            "Add a bulleted list to a slide. Reuses the slide layout's body "
            "placeholder when present; otherwise creates a new text box at the "
            "given position (left/top/width/height required in that case)."
        ),
        parameters={
            "type": "object",
            "properties": {
                "items": {
                    "type": "array",
                    "items": {"type": "string"},
                    "description": "One string per bullet.",
                },
                "slide_index": {
                    "type": "integer",
                    "description": "Defaults to the active slide.",
                },
                "left": {
                    "type": "number",
                    "description": "Left position in inches (only used without a body placeholder).",
                },
                "top": {
                    "type": "number",
                    "description": "Top position in inches (only used without a body placeholder).",
                },
                "width": {
                    "type": "number",
                    "description": "Width in inches (only used without a body placeholder).",
                },
                "height": {
                    "type": "number",
                    "description": "Height in inches (only used without a body placeholder).",
                },
                "font_size": {"type": "integer", "description": "Font size in points."},
            },
            "required": ["items"],
        },
        handler=add_bullet_list,
    )


def _resolve_alignment(alignment: str):
    resolved = _ALIGNMENTS.get(alignment.lower())
    if resolved is None:
        raise ToolExecutionError(
            f"Unknown alignment '{alignment}'. Supported: {sorted(_ALIGNMENTS)}."
        )
    return resolved


def _resolve_body_text_frame(slide, left, top, width, height):
    """Find a non-title placeholder to reuse, or create a new text box."""
    for shape in slide.placeholders:
        is_title = shape.placeholder_format.idx == 0
        if not is_title and shape.has_text_frame:
            return shape.text_frame

    if None in (left, top, width, height):
        raise ToolExecutionError(
            "This slide's layout has no body placeholder to reuse. "
            "Provide left, top, width, and height to create a new text box."
        )
    text_box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    return text_box.text_frame
