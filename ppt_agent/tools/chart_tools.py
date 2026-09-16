"""Tools for adding native, data-driven charts to a slide.

Charts are created with python-pptx's chart API rather than rendered as
static images, so the result stays editable in PowerPoint (double-click
to edit the underlying data table) instead of being a flattened picture.
"""

from __future__ import annotations

from langchain_core.tools import StructuredTool
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from pydantic import BaseModel, Field

from ppt_agent.exceptions import ToolExecutionError
from ppt_agent.presentation.manager import PresentationManager

_CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "stacked_column": XL_CHART_TYPE.COLUMN_STACKED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE,
    "line_markers": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "area": XL_CHART_TYPE.AREA,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
}


class AddChartInput(BaseModel):
    chart_type: str = Field(description=f"Supported: {sorted(_CHART_TYPES)}.")
    categories: list[str] = Field(description="Labels along the category axis, e.g. quarters or product names.")
    series: dict[str, list[float]] = Field(
        description="Mapping of series name to a list of numeric values, one per category."
    )
    slide_index: int | None = Field(default=None, description="Defaults to the active slide.")
    left: float = Field(default=1.0, description="Left position in inches.")
    top: float = Field(default=1.5, description="Top position in inches.")
    width: float = Field(default=8.0, description="Width in inches.")
    height: float = Field(default=4.5, description="Height in inches.")
    title: str | None = Field(default=None, description="Optional chart title.")


def build_tools(manager: PresentationManager) -> list[StructuredTool]:
    """Build chart-related tools bound to ``manager``."""

    def add_chart(
        chart_type: str,
        categories: list[str],
        series: dict[str, list[float]],
        slide_index: int | None = None,
        left: float = 1.0,
        top: float = 1.5,
        width: float = 8.0,
        height: float = 4.5,
        title: str | None = None,
    ) -> str:
        xl_chart_type = _CHART_TYPES.get(chart_type.lower())
        if xl_chart_type is None:
            raise ToolExecutionError(f"Unknown chart_type '{chart_type}'. Supported: {sorted(_CHART_TYPES)}.")
        if not categories:
            raise ToolExecutionError("categories must be a non-empty list.")
        if not series:
            raise ToolExecutionError("series must be a non-empty mapping of series name to values.")
        for series_name, values in series.items():
            if len(values) != len(categories):
                raise ToolExecutionError(
                    f"Series '{series_name}' has {len(values)} value(s) but there are "
                    f"{len(categories)} categories; each series must match the category count."
                )

        slide = manager.get_slide(slide_index)
        chart_data = CategoryChartData()
        chart_data.categories = categories
        for series_name, values in series.items():
            chart_data.add_series(series_name, values)

        graphic_frame = slide.shapes.add_chart(
            xl_chart_type, Inches(left), Inches(top), Inches(width), Inches(height), chart_data
        )
        chart = graphic_frame.chart
        chart.has_title = bool(title)
        if title:
            chart.chart_title.text_frame.text = title

        return (
            f"Added {chart_type} chart with {len(series)} series to "
            f"slide {manager.resolve_slide_index(slide_index)}."
        )

    return [
        StructuredTool.from_function(
            func=add_chart,
            name="add_chart",
            description=(
                "Add a native, editable chart to a slide from category/series data. "
                f"Supported chart_type values: {sorted(_CHART_TYPES)}."
            ),
            args_schema=AddChartInput,
        ),
    ]
